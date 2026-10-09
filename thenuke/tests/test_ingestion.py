"""Comprehensive Test Suite for thenuke Ingestion & Extraction Engine.

Covers:
- UnifiedCorpus data models, JSON serialization, and topic extraction.
- YouTube ingestion: URL validation, VTT parsing, 3-line rolling deduplication, chapter mapping, subprocess orchestration.
- Web crawler: DOM noise pruning, HTML-to-Markdown conversion, code block language extraction, table formatting, SPA detection.
- Local file extraction: Real PyMuPDF PDF parsing (text, tables, images), python-pptx extraction (titles, bullets, tables, speaker notes), Apple Vision OCR, media transcription commands.
- End-to-end unified ingestion pipeline runner.
"""

from __future__ import annotations

import json
import os
import subprocess
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation

from scripts.ingestion import (
    Chapter,
    CodeBlock,
    ExtractedImage,
    IngestionSource,
    UnifiedCorpus,
    align_cues_to_chapters,
    clean_vtt_content,
    crawl_url,
    deduplicate_vtt_cues,
    extract_directory,
    extract_image,
    extract_local_file,
    extract_pdf,
    extract_playlist_entries,
    extract_pptx,
    extract_video_id,
    extract_video_metadata,
    find_chrome_binary,
    html_to_markdown_and_code,
    ingest_youtube_url,
    ingest_youtube_video,
    is_playlist_url,
    is_youtube_url,
    parse_vtt_cues,
    perform_apple_vision_ocr,
    run_ingestion,
)
from scripts.ingestion.file_extract import detect_code_blocks, extract_media
from scripts.ingestion.web_crawler import is_spa_shell, split_markdown_into_chapters
from scripts.ingestion.yt_ingest import format_seconds_timestamp, parse_timestamp_seconds, strip_vtt_markup


# ===========================================================================
# 1. Unified Corpus Data Model & Schema Tests
# ===========================================================================

class TestUnifiedCorpusSchema:
    def test_chapter_dataclass_and_serialization(self):
        chap = Chapter(
            chapter_index=1,
            title="Introduction to V8",
            start_time=0.0,
            end_time=120.5,
            text="V8 compiles JavaScript directly to native machine code.",
        )
        d = chap.to_dict()
        assert d["chapter_index"] == 1
        assert d["title"] == "Introduction to V8"
        assert d["start_time"] == 0.0
        assert d["end_time"] == 120.5
        assert "V8 compiles" in d["text"]

        # Roundtrip
        reconstructed = Chapter.from_dict(d)
        assert reconstructed.chapter_index == chap.chapter_index
        assert reconstructed.title == chap.title
        assert reconstructed.end_time == chap.end_time

    def test_code_block_dataclass(self):
        cb = CodeBlock(language="javascript", code="const answer = 42;")
        d = cb.to_dict()
        assert d["language"] == "javascript"
        assert d["code"] == "const answer = 42;"

        reconstructed = CodeBlock.from_dict(d)
        assert reconstructed.language == "javascript"
        assert reconstructed.code == "const answer = 42;"

    def test_extracted_image_dataclass(self):
        img = ExtractedImage(path="assets/img1.png", caption="Heap graph", ocr_text="Object Pointer")
        d = img.to_dict()
        assert d["path"] == "assets/img1.png"
        assert d["caption"] == "Heap graph"
        assert d["ocr_text"] == "Object Pointer"

        reconstructed = ExtractedImage.from_dict(d)
        assert reconstructed.ocr_text == "Object Pointer"

    def test_ingestion_source_helper_methods(self):
        src = IngestionSource(
            source_type="youtube",
            identifier="https://youtube.com/watch?v=123",
            title="Async JS",
            metadata={"uploader": "Acme", "duration": 500},
        )
        c1 = src.add_chapter(title="Event Loop", text="Microtask queue", start_time=0.0, end_time=100.0)
        cb1 = src.add_code_block(code="queueMicrotask(() => {});", language="javascript")
        im1 = src.add_image(path="assets/loop.png", caption="Tick Diagram")

        assert len(src.chapters) == 1
        assert src.chapters[0].title == "Event Loop"
        assert len(src.extracted_code_blocks) == 1
        assert len(src.extracted_images) == 1

        d = src.to_dict()
        assert d["source_type"] == "youtube"
        assert d["metadata"]["duration"] == 500
        assert len(d["chapters"]) == 1

        reconstructed = IngestionSource.from_dict(d)
        assert reconstructed.title == "Async JS"
        assert len(reconstructed.chapters) == 1

    def test_unified_corpus_full_contract_compliance(self):
        with tempfile.TemporaryDirectory() as td:
            corpus_path = Path(td) / "nuke_ingestion_corpus.json"
            corpus = UnifiedCorpus()

            src = IngestionSource(
                source_type="web",
                identifier="https://v8.dev",
                title="V8 Engine Internals",
                metadata={"uploader": "", "duration": 0, "file_type": "web_article"},
            )
            src.add_chapter("Ignition Bytecode", "Ignition is a register-based interpreter.", 0.0, 0.0)
            src.add_code_block("function test() {}", "javascript")
            src.add_image("assets/v8.png", "V8 Pipeline")
            corpus.add_source(src)

            # Check topic extraction
            topics = corpus.extract_topics()
            assert any("v8" in t.lower() or "engine" in t.lower() or "ignition" in t.lower() for t in topics)

            # Save and verify JSON schema structure
            corpus.save(corpus_path)
            assert corpus_path.exists()

            with open(corpus_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            # Contract checks conforming to PROJECT.md § 4.1
            assert "corpus_id" in data
            assert "timestamp" in data
            assert "sources" in data
            assert "aggregated_topics" in data
            assert isinstance(data["sources"], list)
            assert len(data["sources"]) == 1

            s0 = data["sources"][0]
            assert s0["source_type"] == "web"
            assert s0["identifier"] == "https://v8.dev"
            assert "chapters" in s0
            assert "extracted_code_blocks" in s0
            assert "extracted_images" in s0

            # Reload test
            loaded = UnifiedCorpus.load(corpus_path)
            assert loaded.corpus_id == corpus.corpus_id
            assert len(loaded.sources) == 1
            assert loaded.sources[0].title == "V8 Engine Internals"


# ===========================================================================
# 2. YouTube Ingestion & VTT Deduplication Tests
# ===========================================================================

class TestYouTubeIngestion:
    def test_is_youtube_url(self):
        assert is_youtube_url("https://www.youtube.com/watch?v=dQw4w9WgXcQ")
        assert is_youtube_url("https://youtu.be/dQw4w9WgXcQ")
        assert is_youtube_url("https://m.youtube.com/watch?v=dQw4w9WgXcQ")
        assert is_youtube_url("https://music.youtube.com/watch?v=dQw4w9WgXcQ")
        assert not is_youtube_url("https://vimeo.com/123456")
        assert not is_youtube_url("https://example.com")
        assert not is_youtube_url("")

    def test_is_playlist_url(self):
        assert is_playlist_url("https://www.youtube.com/playlist?list=PLbpi6ZahtOH6Blw3RGY36fl262iNq4O_2")
        assert is_playlist_url("https://www.youtube.com/watch?v=abc12345678&list=PLbpi6ZahtOH6Blw3RGY36fl262iNq4O_2")
        assert not is_playlist_url("https://www.youtube.com/watch?v=abc12345678")

    def test_extract_video_id(self):
        assert extract_video_id("https://www.youtube.com/watch?v=dQw4w9WgXcQ") == "dQw4w9WgXcQ"
        assert extract_video_id("https://youtu.be/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
        assert extract_video_id("https://www.youtube.com/embed/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
        assert extract_video_id("https://www.youtube.com/shorts/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
        assert extract_video_id("https://example.com") is None

    def test_timestamp_conversions(self):
        assert parse_timestamp_seconds("00:01:23.456") == pytest.approx(83.456)
        assert parse_timestamp_seconds("01:23:45,678") == pytest.approx(5025.678)
        assert parse_timestamp_seconds("12.5") == pytest.approx(12.5)
        assert format_seconds_timestamp(83.456) == "00:01:23"
        assert format_seconds_timestamp(3665.0) == "01:01:05"

    def test_strip_vtt_markup(self):
        raw = "<00:00:01.000><c.colorE5E5E5>hello</c> <v Speaker 1>world &amp; friends</v>"
        cleaned = strip_vtt_markup(raw)
        assert cleaned == "hello world & friends"

    def test_vtt_rolling_window_deduplication_algorithm(self):
        """Simulates authentic YouTube auto-generated captions with 3-line rolling sliding window."""
        vtt_data = """WEBVTT
Kind: captions
Language: en

00:00:01.000 --> 00:00:03.000
Welcome everyone
to this session
on V8 internals

00:00:02.500 --> 00:00:04.500
to this session
on V8 internals
we will discuss the heap

00:00:04.000 --> 00:00:06.000
on V8 internals
we will discuss the heap
and the event loop
"""
        cues = parse_vtt_cues(vtt_data)
        assert len(cues) == 3

        deduped = deduplicate_vtt_cues(cues)
        full_text = clean_vtt_content(vtt_data)

        # Confirm the repeating lines were not duplicated 3x
        assert full_text.count("Welcome everyone") == 1
        assert full_text.count("to this session") == 1
        assert full_text.count("on V8 internals") == 1
        assert full_text.count("we will discuss the heap") == 1
        assert full_text.count("and the event loop") == 1

    def test_chapter_alignment_with_metadata_chapters(self):
        cues = [
            {"start": 10.0, "end": 20.0, "text": "Part one content"},
            {"start": 70.0, "end": 80.0, "text": "Part two content"},
        ]
        raw_chapters = [
            {"title": "Introduction", "start_time": 0.0, "end_time": 60.0},
            {"title": "Architecture", "start_time": 60.0, "end_time": 120.0},
        ]
        chapters = align_cues_to_chapters(cues, raw_chapters, total_duration=120.0)
        assert len(chapters) == 2
        assert chapters[0].title == "Introduction"
        assert chapters[0].text == "Part one content"
        assert chapters[1].title == "Architecture"
        assert chapters[1].text == "Part two content"

    def test_chapter_alignment_without_chapters_creates_single_or_segments(self):
        cues = [
            {"start": 5.0, "end": 15.0, "text": "Hello world from transcript"},
        ]
        chapters = align_cues_to_chapters(cues, [], total_duration=300.0, video_title="Sample Video")
        assert len(chapters) == 1
        assert "Sample Video" in chapters[0].title
        assert "Hello world" in chapters[0].text

    @patch("scripts.ingestion.yt_ingest.run_cmd")
    def test_extract_video_metadata_mocked(self, mock_run_cmd):
        mock_meta = {
            "id": "dQw4w9WgXcQ",
            "title": "Mocked YouTube Video",
            "duration": 212,
            "uploader": "Test Channel",
            "chapters": [{"title": "Ch 1", "start_time": 0.0, "end_time": 212.0}],
        }
        mock_run_cmd.return_value = (0, json.dumps(mock_meta) + "\n", "")

        meta = extract_video_metadata("https://www.youtube.com/watch?v=dQw4w9WgXcQ")
        assert meta["title"] == "Mocked YouTube Video"
        assert meta["duration"] == 212

    @patch("scripts.ingestion.yt_ingest.run_cmd")
    def test_extract_playlist_entries_mocked(self, mock_run_cmd):
        lines = [
            json.dumps({"id": "vid1", "title": "Video 1", "url": "https://youtube.com/watch?v=vid1"}),
            json.dumps({"id": "vid2", "title": "Video 2", "url": "https://youtube.com/watch?v=vid2"}),
        ]
        mock_run_cmd.return_value = (0, "\n".join(lines), "")

        entries = extract_playlist_entries("https://www.youtube.com/playlist?list=PL123")
        assert len(entries) == 2
        assert entries[0]["id"] == "vid1"
        assert entries[1]["title"] == "Video 2"

    @patch("scripts.ingestion.yt_ingest.extract_video_metadata")
    @patch("scripts.ingestion.yt_ingest.download_subtitles")
    def test_ingest_youtube_video_end_to_end(self, mock_down_subs, mock_meta):
        mock_meta.return_value = {
            "id": "vid_xyz",
            "title": "Event Loop Deep Dive",
            "duration": 600,
            "uploader": "Sparsh Tech",
            "chapters": [{"title": "Phase 1: Macrotasks", "start_time": 0.0, "end_time": 600.0}],
        }

        with tempfile.TemporaryDirectory() as td:
            sub_file = Path(td) / "vid_xyz.en.vtt"
            sub_file.write_text("""WEBVTT\n\n00:00:10.000 --> 00:00:20.000\nMacrotask queue execution details.\n""")
            mock_down_subs.return_value = sub_file

            source = ingest_youtube_video("https://www.youtube.com/watch?v=vid_xyz", output_dir=td)
            assert source.source_type == "youtube"
            assert source.title == "Event Loop Deep Dive"
            assert source.metadata["uploader"] == "Sparsh Tech"
            assert len(source.chapters) == 1
            assert "Macrotask" in source.chapters[0].text


# ===========================================================================
# 3. Web Crawler Tests
# ===========================================================================

class TestWebCrawler:
    def test_find_chrome_binary(self):
        chrome_path = find_chrome_binary()
        assert chrome_path is not None
        assert os.path.exists(chrome_path)

    def test_spa_detection_logic(self):
        assert is_spa_shell("<html><body><div id='root'></div><script src='bundle.js'></script></body></html>")
        assert is_spa_shell("<html><body><div id='app'></div></body></html>")
        assert not is_spa_shell(
            "<html><body><main><h1>Full Documentation Page</h1>"
            "<p>This is a complete static HTML article with sufficient content "
            "explaining V8 compiler optimization pipelines in full detail. "
            "It has hundreds of characters of genuine prose without empty mounting nodes.</p></main></body></html>"
        )

    def test_html_to_markdown_and_code_extraction(self):
        html_input = """
        <!DOCTYPE html>
        <html>
        <head><title>MDN Web Docs: Promises</title></head>
        <body>
            <nav><a href="/">Nav Item</a></nav>
            <div class="sidebar">Sidebar menu</div>
            <main>
                <h1>Promise.withResolvers()</h1>
                <p>The <code>Promise.withResolvers()</code> static method returns an object containing a promise.</p>
                
                <h2>Syntax and Example</h2>
                <pre><code class="language-javascript">
const { promise, resolve, reject } = Promise.withResolvers();
resolve("Success!");
                </code></pre>

                <h2>Browser Compatibility</h2>
                <table>
                    <tr><th>Browser</th><th>Version</th></tr>
                    <tr><td>Chrome</td><td>119</td></tr>
                    <tr><td>Node.js</td><td>22</td></tr>
                </table>
                <blockquote>A standardized helper method in ES2024.</blockquote>
            </main>
            <footer>Footer notes</footer>
        </body>
        </html>
        """
        md, code_blocks, title = html_to_markdown_and_code(html_input)

        assert title == "MDN Web Docs: Promises"
        assert "# Promise.withResolvers()" in md
        assert "## Syntax and Example" in md
        assert "```javascript" in md
        assert "resolve(\"Success!\");" in md
        assert "| Chrome | 119 |" in md
        assert "> A standardized helper method in ES2024." in md
        # Nav and footer must be removed
        assert "Nav Item" not in md
        assert "Footer notes" not in md

        # Code block objects
        assert len(code_blocks) == 1
        assert code_blocks[0].language == "javascript"
        assert "Promise.withResolvers" in code_blocks[0].code

    def test_split_markdown_into_chapters(self):
        md_text = """
# Introduction to Garbage Collection
Generational hypothesis in V8.

## The Scavenger (Young Generation)
Semi-space copying algorithm.

## Mark-Sweep-Compact (Old Generation)
Three phase major GC.
"""
        chapters = split_markdown_into_chapters(md_text)
        assert len(chapters) == 3
        assert "Introduction to Garbage Collection" in chapters[0].title
        assert "Scavenger" in chapters[1].title
        assert "Mark-Sweep-Compact" in chapters[2].title

    @patch("scripts.ingestion.web_crawler.requests.get")
    def test_crawl_url_tier1_mocked(self, mock_get):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = """
        <html><head><title>Test Article</title></head>
        <body><main><h1>Node.js Architecture</h1><p>Node uses libuv for asynchronous I/O operations.</p>
        <pre><code class="language-js">console.log("hello");</code></pre>
        </main></body></html>
        """
        mock_get.return_value = mock_resp

        source = crawl_url("https://example.com/node-arch", capture_screenshot=False)
        assert source.source_type == "web"
        assert source.title == "Test Article"
        assert source.metadata["crawler_tier"] == "tier1_requests"
        assert len(source.chapters) >= 1
        assert len(source.extracted_code_blocks) == 1


# ===========================================================================
# 4. Local File Extractor Tests
# ===========================================================================

class TestLocalFileExtractor:
    def test_detect_code_blocks(self):
        text = """
Here is an example in JavaScript:
```javascript
function add(a, b) {
    return a + b;
}
```
And some text.
"""
        blocks = detect_code_blocks(text)
        assert len(blocks) == 1
        assert blocks[0].language == "javascript"
        assert "return a + b;" in blocks[0].code

    def test_pdf_extraction_with_real_pymupdf_document(self):
        """Constructs an authentic PDF in memory using PyMuPDF and validates full extraction."""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = Path(td) / "sample_v8.pdf"
            doc = pymupdf.open()
            
            # Page 1
            p1 = doc.new_page()
            p1.insert_text((50, 50), "# V8 Execution Pipeline", fontsize=18)
            p1.insert_text((50, 80), "Ignition is the interpreter. TurboFan is the optimizing compiler.", fontsize=12)

            # Insert a table on page 1
            # (Drawing a simple text block representation)
            p1.insert_text((50, 120), "| Phase | Tool |\n| --- | --- |\n| Parse | Scanner |\n| JIT | TurboFan |", fontsize=10)

            # Insert an image on page 2
            p2 = doc.new_page()
            p2.insert_text((50, 50), "# Memory Layout", fontsize=18)
            
            # Create a small image and embed
            img_pil = Image.new("RGB", (80, 80), color=(100, 150, 200))
            img_temp_path = Path(td) / "temp_embed.png"
            img_pil.save(img_temp_path)
            p2.insert_image(pymupdf.Rect(50, 80, 150, 180), filename=str(img_temp_path))

            doc.save(str(pdf_path))
            doc.close()

            # Now run extract_pdf
            source = extract_pdf(pdf_path, output_dir=Path(td))
            assert source.source_type == "local_file"
            assert source.metadata["file_type"] == "pdf"
            assert source.metadata["page_count"] == 2
            assert len(source.chapters) == 2
            assert "V8 Execution Pipeline" in source.chapters[0].text
            assert len(source.extracted_images) >= 1
            assert Path(source.extracted_images[0].path).exists()

    def test_pptx_extraction_with_real_presentation(self):
        """Constructs an authentic PPTX using python-pptx and validates slide, bullets, table, and notes extraction."""
        with tempfile.TemporaryDirectory() as td:
            pptx_path = Path(td) / "sample_deck.pptx"
            prs = Presentation()

            # Slide 1: Title and body bullets
            slide_layout = prs.slide_layouts[1]  # Bullet slide layout
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = "Event Loop Invariants"
            tf = slide.placeholders[1].text_frame
            tf.text = "Rule 1: Microtasks drain before macrotasks."
            p2 = tf.add_paragraph()
            p2.text = "Rule 2: Render steps occur between macrotask ticks."

            # Add speaker notes
            notes = slide.notes_slide.notes_text_frame
            notes.text = "GOTCHA: Starvation occurs if microtasks schedule more microtasks infinitely."

            # Slide 2: Table
            blank_layout = prs.slide_layouts[6]
            slide2 = prs.slides.add_slide(blank_layout)
            table_shape = slide2.shapes.add_table(rows=2, cols=2, left=100, top=100, width=5000000, height=2000000)
            table = table_shape.table
            table.cell(0, 0).text = "Queue"
            table.cell(0, 1).text = "Priority"
            table.cell(1, 0).text = "Microtask"
            table.cell(1, 1).text = "Immediate"

            prs.save(str(pptx_path))

            # Run extract_pptx
            source = extract_pptx(pptx_path, output_dir=Path(td))
            assert source.source_type == "local_file"
            assert source.metadata["file_type"] == "pptx"
            assert source.metadata["slide_count"] == 2
            assert len(source.chapters) == 2

            # Verify slide 1 contents
            ch1_text = source.chapters[0].text
            assert "Event Loop Invariants" in ch1_text
            assert "Microtasks drain before macrotasks" in ch1_text
            assert "[SPEAKER NOTES / GOTCHAS]:" in ch1_text
            assert "Starvation occurs" in ch1_text

            # Verify slide 2 table contents
            ch2_text = source.chapters[1].text
            assert "| Queue | Priority |" in ch2_text
            assert "| Microtask | Immediate |" in ch2_text

    def test_image_extraction_and_apple_vision_ocr(self):
        """Constructs an authentic test image with bold text and tests macOS Apple Vision OCR extraction."""
        with tempfile.TemporaryDirectory() as td:
            img_path = Path(td) / "test_diagram.png"
            img = Image.new("RGB", (400, 120), color=(255, 255, 255))
            draw = ImageDraw.Draw(img)
            draw.text((20, 45), "CLOSURE LEXICAL ENVIRONMENT", fill=(0, 0, 0))
            img.save(img_path)

            source = extract_image(img_path, output_dir=Path(td))
            assert source.source_type == "local_file"
            assert source.metadata["file_type"] == "image"
            assert source.metadata["width"] == 400
            assert source.metadata["height"] == 120
            assert len(source.extracted_images) == 1

            # Check that Apple Vision OCR ran on macOS
            ocr_text = perform_apple_vision_ocr(img_path)
            assert "CLOSURE" in ocr_text or "ENVIRONMENT" in ocr_text

    def test_directory_extraction_heterogeneous(self):
        """Tests recursive directory extraction for mixed documents, images, and text."""
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            
            # File 1: Markdown
            (root / "doc1.md").write_text("# Chapter 1\nOverview text\n```js\nlet a = 1;\n```\n")

            # File 2: Image
            img = Image.new("RGB", (50, 50), color=(10, 20, 30))
            img.save(root / "diagram.png")

            # File 3: Subdirectory with text
            sub = root / "subfolder"
            sub.mkdir()
            (sub / "notes.txt").write_text("Engine gotcha notes")

            sources = extract_directory(root, output_dir=root / "out")
            assert len(sources) >= 3
            identifiers = [s.identifier for s in sources]
            assert any("doc1.md" in i for i in identifiers)
            assert any("diagram.png" in i for i in identifiers)
            assert any("notes.txt" in i for i in identifiers)


# ===========================================================================
# 5. Unified Ingestion Runner End-to-End Integration Tests
# ===========================================================================

class TestUnifiedIngestionRunner:
    def test_run_ingestion_heterogeneous_sources(self):
        with tempfile.TemporaryDirectory() as td:
            work_dir = Path(td)
            out_corpus = work_dir / "nuke_ingestion_corpus.json"

            # Create test local file
            local_doc = work_dir / "architecture.md"
            local_doc.write_text("""
# V8 Heap Architecture
The heap is divided into distinct spaces:
1. New Space (Semi-space allocation)
2. Old Pointer Space
3. Old Data Space
4. Large Object Space
5. Code Space
""")

            # Create test image
            img_path = work_dir / "memory_graph.png"
            im = Image.new("RGB", (200, 100), color=(240, 240, 240))
            draw = ImageDraw.Draw(im)
            draw.text((10, 40), "STACK VS HEAP POINTER", fill=(0, 0, 0))
            im.save(img_path)

            # Run ingestion
            sources_input = [str(local_doc), str(img_path)]
            corpus = run_ingestion(
                sources=sources_input,
                output_corpus_path=out_corpus,
                assets_dir=work_dir / "assets",
            )

            assert len(corpus.sources) == 2
            assert out_corpus.exists()

            # Verify persisted JSON matches interface contract
            loaded = UnifiedCorpus.load(out_corpus)
            assert len(loaded.sources) == 2
            assert loaded.sources[0].title == "architecture"
            assert "New Space" in loaded.sources[0].chapters[0].text
            assert len(loaded.aggregated_topics) > 0


# ===========================================================================
# 6. Edge Cases & Resilience Tests
# ===========================================================================

class TestIngestionEdgeCasesAndResilience:
    def test_invalid_youtube_urls_raise_value_error(self):
        with pytest.raises(ValueError, match="Invalid YouTube URL"):
            ingest_youtube_video("https://vimeo.com/98765")

        with pytest.raises(ValueError, match="Invalid YouTube URL"):
            ingest_youtube_url("not_a_valid_url")

    def test_invalid_web_url_raises_value_error(self):
        with pytest.raises(ValueError, match="Invalid web URL"):
            crawl_url("ftp://example.com/file.txt")

        with pytest.raises(ValueError, match="Invalid web URL"):
            crawl_url("file:///etc/passwd")

    def test_nonexistent_local_file_raises_filenotfound(self):
        with pytest.raises(FileNotFoundError):
            extract_local_file("/non/existent/path/to/file.pdf")

    def test_nonexistent_local_directory_raises_notadirectory(self):
        with pytest.raises(NotADirectoryError):
            extract_directory("/non/existent/directory/path")

    @patch("scripts.ingestion.yt_ingest.run_cmd")
    @patch("scripts.ingestion.yt_ingest.download_subtitles")
    @patch("scripts.ingestion.yt_ingest.transcribe_audio_fallback")
    @patch("scripts.ingestion.yt_ingest.extract_video_metadata")
    def test_youtube_audio_transcription_fallback_when_no_subs(
        self, mock_meta, mock_fallback, mock_down_subs, mock_cmd
    ):
        mock_meta.return_value = {
            "id": "nosub123",
            "title": "Unsubtitled Video",
            "duration": 300,
            "uploader": "AudioOnly",
            "chapters": [],
        }
        mock_down_subs.return_value = None  # No subtitle track
        mock_cmd.return_value = (0, "https://googlevideo.com/stream\n", "")
        mock_fallback.return_value = "This audio was transcribed using Apple Silicon whisper."

        with tempfile.TemporaryDirectory() as td:
            source = ingest_youtube_video("https://www.youtube.com/watch?v=nosub123", output_dir=td)
            assert source.source_type == "youtube"
            assert source.title == "Unsubtitled Video"
            assert len(source.chapters) == 1
            assert "whisper" in source.chapters[0].text

    @patch("scripts.ingestion.file_extract.run_cmd")
    def test_local_media_extraction_mocked(self, mock_cmd):
        with tempfile.TemporaryDirectory() as td:
            video_file = Path(td) / "lecture.mp4"
            video_file.write_bytes(b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00")
            
            # Subtitle check fails, audio ffmpeg succeeds
            mock_cmd.side_effect = [
                (1, "", "no subs"),  # ffmpeg sub check
                (0, "", ""),         # ffmpeg audio conversion
                (0, "Transcribed lecture on event loops.", ""),  # whisper-cli
            ]

            source = extract_media(video_file, output_dir=Path(td))
            assert source.source_type == "local_file"
            assert source.metadata["file_type"] == "mp4"
            assert len(source.chapters) == 1
            assert "lecture on event loops" in source.chapters[0].text

    def test_malformed_vtt_cues_graceful_recovery(self):
        corrupt_vtt = """WEBVTT
NOT A CUE
00:00:01.000 -> INVALID TIMESTAMP
broken line
00:00:05.000 --> 00:00:08.000
Valid cue text here.
"""
        cues = parse_vtt_cues(corrupt_vtt)
        assert len(cues) == 1
        assert cues[0]["start"] == 5.0
        assert cues[0]["text"] == "Valid cue text here."

    def test_unified_corpus_empty_state_and_resilience(self):
        corpus = UnifiedCorpus()
        assert corpus.sources == []
        assert corpus.aggregated_topics == []
        topics = corpus.extract_topics()
        assert topics == []

        d = corpus.to_dict()
        assert d["sources"] == []
        reconstructed = UnifiedCorpus.from_dict(d)
        assert reconstructed.sources == []
